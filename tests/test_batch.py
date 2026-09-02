"""
Tests for the bulk spreadsheet upload: column detection, row validation, the
job lifecycle, and the .xlsx exports. Scrapers are faked — no browser/network.
"""
from __future__ import annotations

import io
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import batch  # noqa: E402
import index  # noqa: E402
from index import create_app  # noqa: E402
from models import Carrier, Status, TrackingEvent, TrackingResult  # noqa: E402

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def make_xlsx(rows: list[list], headers: list[str] | None = ["company", "awb"]) -> bytes:
    wb = Workbook()
    ws = wb.active
    if headers:
        ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _FakeScraper:
    """Returns a delivered result, or raises when `boom` is set."""

    def __init__(self, carrier: Carrier = Carrier.DHL, boom: str | None = None):
        self.carrier = carrier
        self.boom = boom
        self.calls: list[str] = []

    def scrape(self, number: str) -> TrackingResult:
        self.calls.append(number)
        if self.boom:
            raise RuntimeError(self.boom)
        return TrackingResult(
            tracking_number=number,
            carrier=self.carrier,
            status=Status.DELIVERED,
            origin="Leipzig",
            destination="Dublin",
            delivered_at=datetime(2026, 6, 30, 10, 25, tzinfo=timezone.utc),
            events=[
                TrackingEvent(
                    timestamp=datetime(2026, 6, 30, 10, 25, tzinfo=timezone.utc),
                    location="Dublin",
                    description="Delivered",
                    status=Status.DELIVERED,
                )
            ],
            scraped_at=datetime(2026, 6, 30, 11, 0, tzinfo=timezone.utc),
        )


@pytest.fixture()
def client() -> TestClient:
    return TestClient(create_app())


@pytest.fixture(autouse=True)
def fresh_registry(monkeypatch):
    """Each test gets its own job registry so jobs never leak between tests."""
    monkeypatch.setattr(batch, "JOBS", batch.JobRegistry())


def upload(client: TestClient, data: bytes, name: str = "shipments.xlsx", **params):
    return client.post(
        "/api/batch",
        files={"file": (name, data, XLSX_MIME)},
        params=params or None,
    )


def wait_for(client: TestClient, job_id: str, timeout: float = 10.0) -> dict:
    """Poll a job until it leaves the running state."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        body = client.get(f"/api/batch/{job_id}").json()
        if body["state"] in {"completed", "cancelled"}:
            return body
        time.sleep(0.02)
    raise AssertionError(f"job {job_id} did not finish within {timeout}s")


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def test_parses_company_and_awb_columns():
    parsed = batch.parse_workbook(
        "s.xlsx", make_xlsx([["DHL", "1790531772"], ["UPS", "1Z999AA10123456784"]])
    )
    assert [r.carrier for r in parsed.rows] == [Carrier.DHL, Carrier.UPS]
    assert [r.awb for r in parsed.rows] == ["1790531772", "1Z999AA10123456784"]
    assert all(r.valid for r in parsed.rows)


@pytest.mark.parametrize(
    "company,expected",
    [
        ("DHL", Carrier.DHL),
        ("dhl express", Carrier.DHL),
        ("  FedEx  ", Carrier.FEDEX),
        ("Federal Express", Carrier.FEDEX),
        ("United Parcel Service", Carrier.UPS),
        ("ARAMEX", Carrier.ARAMEX),
        ("Blue Dart", None),
        ("", None),
    ],
)
def test_resolve_carrier_aliases(company, expected):
    assert batch.resolve_carrier(company) is expected


@pytest.mark.parametrize(
    "headers",
    [["Company", "AWB"], ["carrier", "AWB No."], ["Courier Name", "Tracking Number"]],
)
def test_header_spellings_are_accepted(headers):
    parsed = batch.parse_workbook("s.xlsx", make_xlsx([["DHL", "123"]], headers))
    assert parsed.rows[0].carrier is Carrier.DHL


def test_header_can_sit_below_a_title_row():
    data = make_xlsx([["company", "awb"], ["DHL", "123"]], headers=["Shipment report"])
    parsed = batch.parse_workbook("s.xlsx", data)
    assert len(parsed.rows) == 1
    assert parsed.rows[0].awb == "123"


def test_numeric_awb_does_not_become_a_float():
    parsed = batch.parse_workbook("s.xlsx", make_xlsx([["DHL", 1790531772]]))
    assert parsed.rows[0].awb == "1790531772"


def test_extra_columns_are_ignored():
    data = make_xlsx(
        [["DHL", "123", "our-ref-9"]], headers=["company", "awb", "reference"]
    )
    parsed = batch.parse_workbook("s.xlsx", data)
    assert parsed.rows[0].awb == "123"


def test_blank_rows_are_skipped_not_reported_as_errors():
    parsed = batch.parse_workbook(
        "s.xlsx", make_xlsx([["DHL", "123"], [None, None], ["UPS", "456"]])
    )
    assert len(parsed.rows) == 2
    assert parsed.skipped_blank == 1


def test_invalid_rows_are_kept_with_an_error():
    parsed = batch.parse_workbook(
        "s.xlsx",
        make_xlsx(
            [
                ["DHL", "123"],
                ["Blue Dart", "999"],       # unknown company
                ["DHL", ""],                # missing awb
                ["", "555"],                # missing company
                ["DHL", "12 34/56"],        # illegal characters
                ["dhl", "123"],             # duplicate of row 2
            ]
        ),
    )
    errors = [r.error for r in parsed.rows]
    assert errors[0] is None
    assert "Unknown company" in errors[1]
    assert errors[2] == "Missing AWB number."
    assert errors[3] == "Missing company."
    assert "letters, digits" in errors[4]
    # A repeat names the row it repeats, so the user knows where the answer is.
    assert errors[5].startswith("Same AWB as row 2")
    assert parsed.rows[5].duplicate is True
    assert parsed.rows[5].duplicate_of == 2
    assert parsed.rows[parsed.rows[5].duplicate_of_index].row_number == 2
    assert len(parsed.valid_rows) == 1


def test_row_numbers_point_at_the_source_line():
    parsed = batch.parse_workbook("s.xlsx", make_xlsx([["DHL", "1"], ["UPS", "2"]]))
    # Header is line 1, so the first data row is line 2.
    assert [r.row_number for r in parsed.rows] == [2, 3]


def test_csv_upload_is_supported():
    csv_bytes = b"company,awb\nDHL,1790531772\nUPS,1Z999AA10123456784\n"
    parsed = batch.parse_workbook("s.csv", csv_bytes)
    assert [r.carrier for r in parsed.rows] == [Carrier.DHL, Carrier.UPS]


def test_semicolon_csv_is_supported():
    parsed = batch.parse_workbook("s.csv", b"company;awb\nDHL;123\nUPS;456\n")
    assert [r.awb for r in parsed.rows] == ["123", "456"]


def test_row_cap_truncates_and_flags():
    rows = [["DHL", str(i)] for i in range(batch.MAX_ROWS + 25)]
    parsed = batch.parse_workbook("s.xlsx", make_xlsx(rows))
    assert len(parsed.rows) == batch.MAX_ROWS
    assert parsed.truncated is True


@pytest.mark.parametrize(
    "filename,data,message",
    [
        ("s.xlsx", b"", "empty"),
        ("s.pdf", b"whatever", "Excel"),
        ("s.xls", b"whatever", "Legacy"),
        ("s.xlsx", make_xlsx([], headers=["name", "value"]), "Couldn't find"),
        ("s.csv", b"company,awb\n", "No data rows"),
    ],
)
def test_unusable_files_raise_upload_error(filename, data, message):
    with pytest.raises(batch.UploadError, match=message):
        batch.parse_workbook(filename, data)


def test_oversized_file_is_rejected():
    with pytest.raises(batch.UploadError, match="too large"):
        batch.parse_workbook("s.xlsx", b"x" * (batch.MAX_UPLOAD_BYTES + 1))


# ---------------------------------------------------------------------------
# Job lifecycle over the API
# ---------------------------------------------------------------------------
def test_upload_tracks_every_valid_row(client: TestClient, monkeypatch):
    dhl, ups = _FakeScraper(Carrier.DHL), _FakeScraper(Carrier.UPS)
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, dhl)
    monkeypatch.setitem(index.SCRAPERS, Carrier.UPS, ups)

    r = upload(client, make_xlsx([["DHL", "1790531772"], ["UPS", "1Z9999"]]))
    assert r.status_code == 202
    job = r.json()
    assert job["counts"]["total"] == 2

    done = wait_for(client, job["id"])
    assert done["state"] == "completed"
    assert done["counts"]["done"] == 2
    assert done["progress"] == 1.0
    assert dhl.calls == ["1790531772"]
    assert ups.calls == ["1Z9999"]
    assert done["rows"][0]["result"]["status"] == "delivered"
    assert done["rows"][0]["result"]["destination"] == "Dublin"


def test_failed_row_does_not_stop_the_rest(client: TestClient, monkeypatch):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper(boom="blocked"))
    monkeypatch.setitem(index.SCRAPERS, Carrier.UPS, _FakeScraper(Carrier.UPS))

    job = upload(client, make_xlsx([["DHL", "111"], ["UPS", "222"]])).json()
    done = wait_for(client, job["id"])

    by_awb = {r["awb"]: r for r in done["rows"]}
    assert by_awb["111"]["state"] == "failed"
    assert "blocked" in by_awb["111"]["error"]
    assert by_awb["222"]["state"] == "done"
    assert done["counts"] == {**done["counts"], "done": 1, "failed": 1}


def test_invalid_rows_are_skipped_without_scraping(client: TestClient, monkeypatch):
    dhl = _FakeScraper(Carrier.DHL)
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, dhl)

    job = upload(client, make_xlsx([["Blue Dart", "999"], ["DHL", "111"]])).json()
    done = wait_for(client, job["id"])

    assert dhl.calls == ["111"]
    skipped = next(r for r in done["rows"] if r["awb"] == "999")
    assert skipped["state"] == "skipped"
    assert "Unknown company" in skipped["error"]
    assert any("couldn't be tracked" in n for n in done["notes"])


def test_upload_with_no_trackable_rows_completes_immediately(
    client: TestClient, monkeypatch
):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper())
    job = upload(client, make_xlsx([["Blue Dart", "999"]])).json()
    assert job["state"] == "completed"
    assert job["counts"]["skipped"] == 1


def test_bad_file_returns_400(client: TestClient):
    r = upload(client, b"not a spreadsheet", name="notes.pdf")
    assert r.status_code == 400
    assert "Excel" in r.json()["detail"]


def test_unknown_job_returns_404(client: TestClient):
    assert client.get("/api/batch/deadbeef").status_code == 404
    assert client.post("/api/batch/deadbeef/cancel").status_code == 404
    assert client.get("/api/batch/deadbeef/export.xlsx").status_code == 404


def test_list_jobs_is_newest_first_and_row_free(client: TestClient, monkeypatch):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper())
    first = upload(client, make_xlsx([["DHL", "111"]]), name="first.xlsx").json()
    second = upload(client, make_xlsx([["DHL", "222"]]), name="second.xlsx").json()
    wait_for(client, first["id"])
    wait_for(client, second["id"])

    listed = client.get("/api/batch").json()
    assert [j["id"] for j in listed] == [second["id"], first["id"]]
    assert listed[0]["filename"] == "second.xlsx"
    assert listed[0]["rows"] == []


def test_cancel_skips_queued_rows(client: TestClient, monkeypatch):
    class _Slow(_FakeScraper):
        def scrape(self, number: str) -> TrackingResult:
            time.sleep(0.25)
            return super().scrape(number)

    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _Slow())
    rows = [["DHL", str(i)] for i in range(8)]
    job = upload(client, make_xlsx(rows), concurrency=1).json()

    cancelled = client.post(f"/api/batch/{job['id']}/cancel").json()
    assert cancelled["state"] == "cancelled"

    final = wait_for(client, job["id"])
    assert final["state"] == "cancelled"
    assert final["counts"]["skipped"] >= 1


# ---------------------------------------------------------------------------
# Spreadsheet output
# ---------------------------------------------------------------------------
def test_results_export_has_a_row_per_shipment(client: TestClient, monkeypatch):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper())
    job = upload(client, make_xlsx([["DHL", "111"], ["DHL", "222"]])).json()
    wait_for(client, job["id"])

    r = client.get(f"/api/batch/{job['id']}/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX_MIME
    assert "shipments-results.xlsx" in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Tracking results"]
    header = [c.value for c in ws[1]]
    assert header[:5] == ["Row", "Company", "AWB", "Carrier", "Status"]
    assert ws.max_row == 3  # header + 2 shipments
    assert [ws.cell(row=r_, column=3).value for r_ in (2, 3)] == ["111", "222"]
    assert ws.cell(row=2, column=5).value == "Delivered"
    # Every scan lands on the history sheet.
    assert wb["Event history"].max_row == 3


def test_template_download_is_a_usable_upload(client: TestClient):
    r = client.get("/api/batch/template.xlsx")
    assert r.status_code == 200
    assert "tracking-template.xlsx" in r.headers["content-disposition"]

    parsed = batch.parse_workbook("tracking-template.xlsx", r.content)
    assert {row.carrier for row in parsed.rows} == set(Carrier)


def test_export_of_a_failed_row_carries_the_error(client: TestClient, monkeypatch):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper(boom="wall"))
    job = upload(client, make_xlsx([["DHL", "111"]])).json()
    wait_for(client, job["id"])

    wb = load_workbook(io.BytesIO(client.get(f"/api/batch/{job['id']}/export.xlsx").content))
    ws = wb["Tracking results"]
    assert ws.cell(row=2, column=6).value == "failed"
    # No carrier status is claimed for a row we never actually read.
    assert ws.cell(row=2, column=5).value == "—"
    assert "wall" in ws.cell(row=2, column=ws.max_column).value


def test_duplicates_are_their_own_state_not_a_problem(client: TestClient, monkeypatch):
    """A repeated AWB is informational — real files repeat a waybill per line."""
    dhl = _FakeScraper(Carrier.DHL)
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, dhl)

    job = upload(
        client,
        make_xlsx([["DHL", "111"], ["DHL", "111"], ["Blue Dart", "999"]]),
    ).json()
    done = wait_for(client, job["id"])

    by_awb = {(r["awb"], r["state"]) for r in done["rows"]}
    assert ("111", "done") in by_awb
    assert ("111", "duplicate") in by_awb
    assert ("999", "skipped") in by_awb

    # Tracked once, and the duplicate is counted apart from real problems.
    assert dhl.calls == ["111"]
    assert done["counts"]["duplicate"] == 1
    assert done["counts"]["skipped"] == 1
    assert done["counts"]["finished"] == 3
    assert any("repeat an AWB" in n for n in done["notes"])


def test_a_duplicate_carries_the_result_of_the_row_it_repeats(
    client: TestClient, monkeypatch
):
    """Every row shows the shipment's data, and says where it came from."""
    dhl = _FakeScraper(Carrier.DHL)
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, dhl)

    job = upload(
        client, make_xlsx([["DHL", "111"], ["DHL", "111"], ["dhl", "111"]])
    ).json()
    done = wait_for(client, job["id"])

    source = next(r for r in done["rows"] if r["state"] == "done")
    dupes = [r for r in done["rows"] if r["state"] == "duplicate"]

    assert dhl.calls == ["111"], "still exactly one lookup"
    assert len(dupes) == 2
    for dup in dupes:
        assert dup["duplicate_of"] == source["row_number"]
        assert dup["result"] == source["result"], "the same shipment, same data"
        assert dup["status"] == source["status"]
        assert f"row {source['row_number']}" in dup["error"]

    # One parcel, however many lines quote it.
    assert done["counts"]["delivered"] == 1


def test_a_duplicate_of_a_failed_row_says_so(client: TestClient, monkeypatch):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper(boom="wall"))
    job = upload(client, make_xlsx([["DHL", "111"], ["DHL", "111"]])).json()
    done = wait_for(client, job["id"])

    dup = next(r for r in done["rows"] if r["state"] == "duplicate")
    assert "couldn't be tracked" in dup["error"]
    assert "wall" in dup["error"]


def test_the_export_names_the_duplicated_row(client: TestClient, monkeypatch):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper(Carrier.DHL))
    job = upload(client, make_xlsx([["DHL", "111"], ["DHL", "111"]])).json()
    wait_for(client, job["id"])

    wb = load_workbook(
        io.BytesIO(client.get(f"/api/batch/{job['id']}/export.xlsx").content)
    )
    ws = wb["Tracking results"]
    header = [c.value for c in ws[1]]
    col = header.index("Duplicate of row") + 1
    status = header.index("Status") + 1

    assert ws.cell(row=2, column=col).value is None, "the source row repeats nothing"
    assert ws.cell(row=3, column=col).value == 2, "row 3 repeats file row 2"
    # The duplicate is not a data-less row any more.
    assert ws.cell(row=3, column=status).value == "Delivered"
    # ...but one shipment's history is listed once.
    assert wb["Event history"].max_row == 2


def test_a_duplicate_row_keeps_its_carrier_for_the_ui(client: TestClient, monkeypatch):
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _FakeScraper(Carrier.DHL))
    job = upload(client, make_xlsx([["DHL", "111"], ["DHL Express", "111"]])).json()
    done = wait_for(client, job["id"])

    dup = next(r for r in done["rows"] if r["state"] == "duplicate")
    assert dup["carrier"] == "dhl", "the badge should still show which carrier it was"


# ---------------------------------------------------------------------------
# Work planning — spare workers must land on the carrier that has the rows
# ---------------------------------------------------------------------------
def _unit_sizes(groups, concurrency):
    return sorted(len(rows) for _, rows in batch._plan_units(groups, concurrency))


def test_one_unit_per_carrier_when_concurrency_is_one():
    groups = {"ups": [1] * 180, "dhl": [1] * 3}
    units = batch._plan_units(groups, 1)
    assert [c for c, _ in units] == ["ups", "dhl"]
    assert _unit_sizes(groups, 1) == [3, 180]


def test_the_biggest_carrier_is_split_across_spare_workers():
    """180 UPS + 3 DHL used to leave 3 workers idle behind one 180-row browser."""
    groups = {"ups": [1] * 180, "dhl": [1] * 3}
    units = batch._plan_units(groups, 4)
    assert len(units) == 4
    assert [c for c, _ in units].count("ups") == 3
    assert _unit_sizes(groups, 4) == [3, 45, 45, 90]
    assert sum(len(rows) for _, rows in units) == 183, "no row lost or duplicated"


def test_every_row_is_planned_exactly_once():
    groups = {"ups": list(range(7)), "aramex": list(range(5))}
    units = batch._plan_units(groups, 6)
    planned = [row for _, rows in units for row in rows]
    assert sorted(planned) == sorted(list(range(7)) + list(range(5)))
    for carrier, rows in units:
        assert rows, "an empty unit would open a browser for nothing"


def test_planning_stops_when_there_is_nothing_left_to_split():
    groups = {"dhl": [1]}
    assert len(batch._plan_units(groups, 8)) == 1


# ---------------------------------------------------------------------------
# Browser reuse across a bulk run
# ---------------------------------------------------------------------------
class _SessionScraper:
    """A fake whose session() records how many browsers the run cost."""

    def __init__(self, carrier: Carrier):
        self.carrier = carrier
        self.sessions = 0
        self.calls: list[str] = []

    def session(self, max_lookups=None):
        scraper = self

        class _S:
            def __enter__(self):
                scraper.sessions += 1
                return self

            def __exit__(self, *exc):
                return False

            def track(self, number):
                scraper.calls.append(number)
                return TrackingResult(
                    tracking_number=number,
                    carrier=scraper.carrier,
                    status=Status.DELIVERED,
                )

        return _S()


def test_a_carriers_rows_share_one_browser(client: TestClient, monkeypatch):
    """20 UPS rows must cost 1 browser, not 20."""
    ups = _SessionScraper(Carrier.UPS)
    monkeypatch.setitem(index.SCRAPERS, Carrier.UPS, ups)

    rows = [["UPS", f"1Z{i:06d}"] for i in range(20)]
    job = upload(client, make_xlsx(rows), concurrency=1).json()
    done = wait_for(client, job["id"])

    assert done["counts"]["done"] == 20
    assert len(ups.calls) == 20
    assert ups.sessions == 1, "each row opened its own browser again"


def test_concurrency_buys_browsers_per_worker_not_per_row(
    client: TestClient, monkeypatch
):
    """The point of the split: 4 workers on one carrier, still not 20 browsers."""
    ups = _SessionScraper(Carrier.UPS)
    monkeypatch.setitem(index.SCRAPERS, Carrier.UPS, ups)

    rows = [["UPS", f"1Z{i:06d}"] for i in range(20)]
    job = upload(client, make_xlsx(rows), concurrency=4).json()
    done = wait_for(client, job["id"])

    assert done["counts"]["done"] == 20
    assert sorted(ups.calls) == sorted(f"1Z{i:06d}" for i in range(20))
    assert ups.sessions == 4, "the spare workers should each get a share"


def test_each_carrier_gets_its_own_session(client: TestClient, monkeypatch):
    ups, dhl = _SessionScraper(Carrier.UPS), _SessionScraper(Carrier.DHL)
    monkeypatch.setitem(index.SCRAPERS, Carrier.UPS, ups)
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, dhl)

    # Interleaved in the file — grouping must still collapse them to one each.
    rows = [["UPS", "1Z1"], ["DHL", "D1"], ["UPS", "1Z2"], ["DHL", "D2"]]
    job = upload(client, make_xlsx(rows), concurrency=1).json()
    done = wait_for(client, job["id"])

    assert done["counts"]["done"] == 4
    assert ups.sessions == 1 and dhl.sessions == 1
    assert ups.calls == ["1Z1", "1Z2"]
    assert dhl.calls == ["D1", "D2"]


def test_a_session_that_cannot_open_fails_its_rows_not_the_job(
    client: TestClient, monkeypatch
):
    """Rows behind a dead session must not hang in 'queued' forever."""

    class _Broken:
        def session(self, max_lookups=None):
            raise RuntimeError("Message: invalid session id Stacktrace: #0 0xdead")

    monkeypatch.setitem(index.SCRAPERS, Carrier.UPS, _Broken())
    monkeypatch.setitem(index.SCRAPERS, Carrier.DHL, _SessionScraper(Carrier.DHL))

    job = upload(client, make_xlsx([["UPS", "1Z1"], ["UPS", "1Z2"], ["DHL", "D1"]])).json()
    done = wait_for(client, job["id"])

    assert done["state"] == "completed"
    ups_rows = [r for r in done["rows"] if r["carrier"] == "ups"]
    assert all(r["state"] == "failed" for r in ups_rows)
    assert all("0x" not in (r["error"] or "") for r in ups_rows), "raw stack trace leaked"
    # The other carrier is unaffected.
    assert next(r for r in done["rows"] if r["carrier"] == "dhl")["state"] == "done"
